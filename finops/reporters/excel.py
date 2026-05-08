from __future__ import annotations
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from finops.reporters.models import Report
from finops.models import ResourceCost, Severity

_SEVERITY_COLORS = {
    "CRITICAL": "FF0000",
    "HIGH": "FF8C00",
    "MEDIUM": "FFD700",
    "INFO": "C0C0C0",
}

_TAB_COLORS = {
    "Summary": "2E75B6",
    "Findings": "C00000",
    "Resource Evolution": "375623",
    "Costs by Resource Group": "4A7B5A",
    "Marketplace": "7030A0",
    "Invoices": "595959",
    "AI Insights": "8B4513",
}

_HEADER_BG = "1F5C6B"
_USD_FORMAT = "#,##0.00"
_PCT_FORMAT = "0.00%"
_MAX_COL_WIDTH = 62


def _all_months(report: Report) -> list[str]:
    months: set[str] = set()
    for sub in report.subscriptions:
        if sub.skipped:
            continue
        for c in sub.costs:
            for day in c.daily_costs:
                months.add(day[:7])
    return sorted(months)


def _parse_service_name(cost: ResourceCost) -> str:
    basename = cost.resource_id.split("/")[-1]
    parts = basename.split("-")
    name_parts = []
    for p in parts:
        if len(p) >= 8 and all(c in "0123456789abcdef" for c in p.lower()):
            break
        name_parts.append(p)
    name = "-".join(name_parts).strip("-") if name_parts else ""
    return name or cost.service_name or cost.resource_type.split("/")[-1]


def _style_header(ws: Worksheet, tab_color: str) -> None:
    bg = PatternFill(start_color=_HEADER_BG, end_color=_HEADER_BG, fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = bg
        cell.font = font
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _set_column_widths(ws: Worksheet) -> None:
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, _MAX_COL_WIDTH)


def _apply_usd(ws: Worksheet, col_indices: list[int]) -> None:
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for idx in col_indices:
            row[idx - 1].number_format = _USD_FORMAT


def _finalize(ws: Worksheet, usd_cols: list[int], pct_cols: list[int] | None = None) -> None:
    _set_column_widths(ws)
    _apply_usd(ws, usd_cols)
    if pct_cols:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for idx in pct_cols:
                row[idx - 1].number_format = _PCT_FORMAT


class ExcelReporter:
    def render(self, report: Report) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)
        months = _all_months(report)

        self._summary(wb, report)
        self._findings(wb, report)
        self._resource_evolution(wb, report, months)
        self._costs_by_rg(wb, report, months)
        self._marketplace(wb, report, months)
        self._invoices(wb, report)
        self._ai_insights(wb, report)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _summary(self, wb: Workbook, report: Report) -> None:
        ws = wb.create_sheet("Summary")
        ws.append([
            "Subscription", "Total Cost (USD)",
            "CRITICAL", "HIGH", "MEDIUM", "INFO",
            "Est. Savings/mo (USD)", "Skipped",
        ])
        _style_header(ws, _TAB_COLORS["Summary"])
        for sub in report.subscriptions:
            by_sev = {s.name: 0 for s in Severity}
            for f in sub.findings:
                by_sev[f.severity.name] += 1
            ws.append([
                sub.subscription_name,
                sub.total_cost,
                by_sev["CRITICAL"],
                by_sev["HIGH"],
                by_sev["MEDIUM"],
                by_sev["INFO"],
                sum(f.estimated_monthly_savings_usd for f in sub.findings),
                sub.skipped,
            ])
        _finalize(ws, usd_cols=[2, 7])

    def _findings(self, wb: Workbook, report: Report) -> None:
        ws = wb.create_sheet("Findings")
        ws.append([
            "Subscription", "Resource Group", "Resource", "Type",
            "Severity", "Category", "Est. Savings/mo (USD)", "Recommendation",
        ])
        _style_header(ws, _TAB_COLORS["Findings"])
        severity_col = 5
        for sub in report.subscriptions:
            if sub.skipped:
                continue
            for f in sorted(sub.findings, key=lambda x: x.severity, reverse=True):
                ws.append([
                    sub.subscription_name,
                    f.resource_group,
                    f.resource_id.split("/")[-1],
                    f.resource_type,
                    f.severity.name,
                    f.category,
                    f.estimated_monthly_savings_usd,
                    f.recommendation,
                ])
                color = _SEVERITY_COLORS.get(f.severity.name, "FFFFFF")
                ws.cell(ws.max_row, severity_col).fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid"
                )
        _finalize(ws, usd_cols=[7])

    def _resource_evolution(self, wb: Workbook, report: Report, months: list[str]) -> None:
        ws = wb.create_sheet("Resource Evolution")
        ws.append(["Subscription", "Resource Group", "Resource", "Type"] + months + ["Total (USD)"])
        _style_header(ws, _TAB_COLORS["Resource Evolution"])
        total_col = 5 + len(months)
        month_cols = list(range(5, total_col))

        for sub in report.subscriptions:
            if sub.skipped:
                continue
            azure_costs = [c for c in sub.costs if c.publisher_type.lower() != "marketplace"]
            for c in sorted(azure_costs, key=lambda x: x.total_cost, reverse=True):
                monthly = {}
                for day, amt in c.daily_costs.items():
                    m = day[:7]
                    monthly[m] = monthly.get(m, 0.0) + amt
                month_values = [monthly.get(m, 0.0) for m in months]
                ws.append([
                    sub.subscription_name,
                    c.resource_group,
                    c.resource_id.split("/")[-1],
                    c.resource_type.split("/")[-1],
                ] + month_values + [sum(month_values)])
        _finalize(ws, usd_cols=month_cols + [total_col])

    def _costs_by_rg(self, wb: Workbook, report: Report, months: list[str]) -> None:
        ws = wb.create_sheet("Costs by Resource Group")
        ws.append(["Subscription", "Resource Group"] + months + ["Total (USD)"])
        _style_header(ws, _TAB_COLORS["Costs by Resource Group"])
        total_col = 3 + len(months)
        month_cols = list(range(3, total_col))

        for sub in report.subscriptions:
            if sub.skipped:
                continue
            azure_costs = [c for c in sub.costs if c.publisher_type.lower() != "marketplace"]
            rg_map: dict[str, dict[str, float]] = {}
            for c in azure_costs:
                rg = c.resource_group or "(sin grupo)"
                if rg not in rg_map:
                    rg_map[rg] = {}
                for day, amt in c.daily_costs.items():
                    m = day[:7]
                    rg_map[rg][m] = rg_map[rg].get(m, 0.0) + amt
            for rg, monthly in sorted(rg_map.items()):
                month_values = [monthly.get(m, 0.0) for m in months]
                ws.append([sub.subscription_name, rg] + month_values + [sum(month_values)])
        _finalize(ws, usd_cols=month_cols + [total_col])

    def _marketplace(self, wb: Workbook, report: Report, months: list[str]) -> None:
        ws = wb.create_sheet("Marketplace")
        ws.append(
            ["Subscription", "Resource Group", "Service Name", "Resource Type"]
            + months + ["Total (USD)", "% of Sub"]
        )
        _style_header(ws, _TAB_COLORS["Marketplace"])
        total_col = 5 + len(months)
        pct_col = total_col + 1
        month_cols = list(range(5, total_col))

        for sub in report.subscriptions:
            if sub.skipped:
                continue
            sub_total = sub.total_cost or 1.0
            mp_costs = [c for c in sub.costs if c.publisher_type.lower() == "marketplace"]
            for c in sorted(mp_costs, key=lambda x: x.total_cost, reverse=True):
                monthly = {}
                for day, amt in c.daily_costs.items():
                    m = day[:7]
                    monthly[m] = monthly.get(m, 0.0) + amt
                month_values = [monthly.get(m, 0.0) for m in months]
                row_total = sum(month_values)
                ws.append([
                    sub.subscription_name,
                    c.resource_group,
                    _parse_service_name(c),
                    c.resource_type.split("/")[-1],
                ] + month_values + [row_total, row_total / sub_total])
        _finalize(ws, usd_cols=month_cols + [total_col], pct_cols=[pct_col])

    def _ai_insights(self, wb: Workbook, report: Report) -> None:
        ws = wb.create_sheet("AI Insights")
        ws.append(["Subscription", "Category", "Title", "Detail", "Est. Savings/mo (USD)", "Confidence"])
        _style_header(ws, _TAB_COLORS["AI Insights"])
        for sub in report.subscriptions:
            if sub.skipped:
                continue
            for ins in sub.ai_insights:
                ws.append([
                    sub.subscription_name,
                    ins.category,
                    ins.title,
                    ins.detail,
                    ins.estimated_monthly_savings_usd,
                    ins.confidence,
                ])
        _finalize(ws, usd_cols=[5])

    def _invoices(self, wb: Workbook, report: Report) -> None:
        ws = wb.create_sheet("Invoices")
        ws.append(["Subscription", "Billing Period", "Amount Due (USD)", "Currency", "Status"])
        _style_header(ws, _TAB_COLORS["Invoices"])
        for sub in report.subscriptions:
            if sub.skipped:
                continue
            for inv in sub.invoices:
                ws.append([
                    sub.subscription_name,
                    inv.billing_period,
                    inv.amount_due,
                    inv.currency,
                    inv.status,
                ])
        _finalize(ws, usd_cols=[3])
