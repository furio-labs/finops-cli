from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from finops.reporters.models import report_from_json
from finops.reporters.excel import ExcelReporter

FIXTURE = Path(__file__).parent / "fixtures" / "data.json"


@pytest.fixture
def wb():
    report = report_from_json(FIXTURE.read_text())
    raw = ExcelReporter().render(report)
    return openpyxl.load_workbook(BytesIO(raw))


def test_sheet_names_and_order(wb):
    assert wb.sheetnames == [
        "Summary", "Findings", "Resource Evolution",
        "Costs by Resource Group", "Marketplace", "Invoices", "AI Insights",
    ]


# --- Summary ---

def test_summary_headers(wb):
    ws = wb["Summary"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers == [
        "Subscription", "Total Cost (USD)", "CRITICAL", "HIGH", "MEDIUM", "INFO",
        "Est. Savings/mo (USD)", "Skipped",
    ]


def test_summary_one_row_per_subscription(wb):
    ws = wb["Summary"]
    assert ws.max_row == 2  # header + 1 subscription


def test_summary_values(wb):
    ws = wb["Summary"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["B2"].value == pytest.approx(16.0)  # Azure 10.0 + Marketplace 6.0
    assert ws["D2"].value == 1   # HIGH count
    assert ws["H2"].value is False  # not skipped


# --- Findings ---

def test_findings_headers(wb):
    ws = wb["Findings"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers == [
        "Subscription", "Resource Group", "Resource", "Type",
        "Severity", "Category", "Est. Savings/mo (USD)", "Recommendation",
    ]


def test_findings_one_row_per_finding(wb):
    ws = wb["Findings"]
    assert ws.max_row == 2  # header + 1 finding


def test_findings_values(wb):
    ws = wb["Findings"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["E2"].value == "HIGH"
    assert ws["F2"].value == "Untagged"
    assert ws["H2"].value == "Recurso sin etiquetas requeridas: service."


def test_findings_severity_cell_is_colored(wb):
    ws = wb["Findings"]
    severity_cell = ws["E2"]
    assert severity_cell.fill.patternType == "solid"
    # HIGH → orange FF8C00 (stored as ARGB: FFFF8C00)
    assert severity_cell.fill.fgColor.rgb.upper().endswith("FF8C00")


# --- Costs by Resource Group ---

def test_costs_headers_include_dynamic_months(wb):
    ws = wb["Costs by Resource Group"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers[0] == "Subscription"
    assert headers[1] == "Resource Group"
    assert "2026-05" in headers
    assert headers[-1] == "Total (USD)"


def test_costs_one_row_per_rg(wb):
    ws = wb["Costs by Resource Group"]
    assert ws.max_row == 2  # header + 1 RG


def test_costs_values(wb):
    ws = wb["Costs by Resource Group"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["B2"].value == "rg-prod"
    # Azure-only total: 5.0 + 5.0 (marketplace excluded — it has its own sheet)
    assert ws.cell(2, ws.max_column).value == pytest.approx(10.0)


# --- Invoices ---

def test_invoices_headers(wb):
    ws = wb["Invoices"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers == [
        "Subscription", "Billing Period", "Amount Due (USD)", "Currency", "Status",
    ]


def test_invoices_one_row_per_invoice(wb):
    ws = wb["Invoices"]
    assert ws.max_row == 2  # header + 1 invoice


def test_invoices_values(wb):
    ws = wb["Invoices"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["B2"].value == "202605"
    assert ws["C2"].value == pytest.approx(1500.0)
    assert ws["D2"].value == "USD"
    assert ws["E2"].value == "Due"


# --- Resource Evolution ---

def test_resource_evolution_headers(wb):
    ws = wb["Resource Evolution"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers[0] == "Subscription"
    assert headers[1] == "Resource Group"
    assert headers[2] == "Resource"
    assert headers[3] == "Type"
    assert "2026-05" in headers
    assert headers[-1] == "Total (USD)"


def test_resource_evolution_excludes_marketplace(wb):
    ws = wb["Resource Evolution"]
    # Only the Azure VM row, not the Marketplace entry
    assert ws.max_row == 2  # header + 1 Azure resource


def test_resource_evolution_values(wb):
    ws = wb["Resource Evolution"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["B2"].value == "rg-prod"
    assert ws["C2"].value == "vm1"
    # Total for the Azure VM: 5.0 + 5.0
    assert ws.cell(2, ws.max_column).value == pytest.approx(10.0)


# --- Marketplace ---

def test_marketplace_headers(wb):
    ws = wb["Marketplace"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers[0] == "Subscription"
    assert headers[1] == "Resource Group"
    assert headers[2] == "Service Name"
    assert headers[3] == "Resource Type"
    assert "2026-05" in headers
    assert headers[-2] == "Total (USD)"
    assert headers[-1] == "% of Sub"


def test_marketplace_one_row_per_marketplace_resource(wb):
    ws = wb["Marketplace"]
    assert ws.max_row == 2  # header + 1 marketplace resource


def test_marketplace_values(wb):
    ws = wb["Marketplace"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["B2"].value == "rg-prod"
    assert ws["C2"].value == "datadog"
    # Total: 3.0 + 3.0
    assert ws.cell(2, ws.max_column - 1).value == pytest.approx(6.0)
    # % of sub: 6.0 / 16.0 = 0.375
    assert ws.cell(2, ws.max_column).value == pytest.approx(0.375)


# --- Formatting ---

def test_header_row_has_teal_fill(wb):
    ws = wb["Summary"]
    fill = ws["A1"].fill
    assert fill.patternType == "solid"
    assert fill.fgColor.rgb.upper().endswith("1F5C6B")


def test_summary_tab_color(wb):
    ws = wb["Summary"]
    assert ws.sheet_properties.tabColor is not None
    assert ws.sheet_properties.tabColor.rgb.upper().endswith("2E75B6")


def test_usd_columns_have_number_format(wb):
    ws = wb["Summary"]
    # "Total Cost (USD)" is column B
    assert ws["B2"].number_format == "#,##0.00"


# --- AI Insights ---

def test_ai_insights_headers(wb):
    ws = wb["AI Insights"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers == ["Subscription", "Category", "Title", "Detail", "Est. Savings/mo (USD)", "Confidence"]


def test_ai_insights_one_row_per_insight(wb):
    ws = wb["AI Insights"]
    assert ws.max_row == 2  # header + 1 insight from fixture


def test_ai_insights_values(wb):
    ws = wb["AI Insights"]
    assert ws["A2"].value == "Test Subscription"
    assert ws["B2"].value == "Recommendation"
    assert ws["C2"].value == "VMs sobredimensionadas"
    assert ws["E2"].value == pytest.approx(120.0)
    assert ws["F2"].value == "high"
